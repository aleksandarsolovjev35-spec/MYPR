#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Комплексная проверка оформления отчётов (П-41*.docx) по требованиям.

Проверяет: разрывы страниц перед структурными элементами, отсутствие
длинных тире/многоточий/«немецких» кавычек, поля 3,0/1,5/2,0/2,0 см, скрытый
номер на титульном листе, единый шрифт Times New Roman, выравнивание по ширине,
неразрывность строк таблиц, повторение шапки таблицы, ручную нумерацию списков,
отсутствие «####» в тексте, центрированные подписи к рисункам.
"""

import glob
import os
import re
import sys

import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

ok_all = True


def check(cond, msg):
    global ok_all
    if not cond:
        ok_all = False
        print("   FAIL:", msg)
    return cond


def full_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            seen = set()
            for c in row.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                parts.append(c.text)
    return "\n".join(parts)


def main():
    for path in sorted(glob.glob(os.path.join(ROOT, "П-41*.docx"))):
        name = os.path.basename(path)
        doc = docx.Document(path)
        print(f"== {name}")

        # --- разрывы страниц ---
        breaks = []
        for i, p in enumerate(doc.paragraphs):
            if any(b.get(qn("w:type")) == "page" for b in p._p.iter(qn("w:br"))):
                nxt = next((q.text.strip() for q in doc.paragraphs[i + 1:]
                            if q.text.strip()), "")
                breaks.append(nxt)
        joined = " | ".join(b[:45] for b in breaks)
        # титульный лист отделён разрывом: после первого разрыва идёт
        # содержимое (раздел 1 или эпиграф перед ним)
        check(bool(breaks), "нет ни одного разрыва страниц")
        if breaks:
            first = doc.paragraphs
            after_first = next(
                (i for i, p in enumerate(first)
                 if any(b.get(qn("w:type")) == "page"
                        for b in p._p.iter(qn("w:br")))), None)
            head_txt = "\n".join(p.text for p in first[:after_first])
            body_txt = "\n".join(p.text for p in first[after_first:])
            check(("ОТЧЕТ" in head_txt or "СЕМИНАР" in head_txt)
                  and re.search(r"^1\. ", body_txt, re.M) is not None,
                  f"титульный лист не отделён разрывом: {joined}")
        check(any("Список использованных источников" in b for b in breaks),
              f"нет разрыва перед списком источников: {joined}")
        check(any(b.startswith("Приложение") for b in breaks),
              f"нет разрыва перед приложением: {joined}")

        # --- «AI-маркеры» в тексте ---
        txt = full_text(doc)
        check("—" not in txt, f"длинное тире в тексте: {txt.count('—')}")
        check("…" not in txt, f"многоточие в тексте: {txt.count('…')}")
        check("„" not in txt and "“" not in txt, "кавычки „…“ в тексте")
        check("####" not in txt, "в тексте остался символ ####")

        # --- поля и титул ---
        sec = doc.sections[0]
        check(abs(sec.left_margin - 1080135) < 100, f"левое поле {sec.left_margin}")
        check(abs(sec.right_margin - 539750) < 100, f"правое поле {sec.right_margin}")
        check(abs(sec.top_margin - 720090) < 100, f"верхнее поле {sec.top_margin}")
        check(abs(sec.bottom_margin - 720090) < 100, f"нижнее поле {sec.bottom_margin}")
        check(sec.different_first_page_header_footer,
              "номер страницы выводится на титульном листе (нет titlePg)")

        # --- шрифты ---
        fonts = set()
        for r in doc.element.body.iter(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                continue
            rf = rpr.find(qn("w:rFonts"))
            if rf is not None and rf.get(qn("w:ascii")):
                fonts.add(rf.get(qn("w:ascii")))
        check(fonts <= {"Times New Roman"}, f"посторонние шрифты: {fonts}")

        # --- выравнивание основного текста по ширине ---
        # считаем только тело отчёта (после первого заголовка раздела),
        # исключая заголовки (полужирные) и титульный лист
        start = next((i for i, p in enumerate(doc.paragraphs)
                      if p.style.name == "Heading 1"
                      or re.match(r"^1\. \S", p.text.strip())), 0)
        body_jc = {"both": 0, "other": 0}
        for p in doc.paragraphs[start:]:
            if not p.text.strip():
                continue
            if p.runs and all(r.bold for r in p.runs):
                continue  # заголовок
            ppr = p._p.find(qn("w:pPr"))
            jc = ppr.find(qn("w:jc")) if ppr is not None else None
            val = jc.get(qn("w:val")) if jc is not None else "none"
            if val == "both":
                body_jc["both"] += 1
            else:
                body_jc["other"] += 1
        check(body_jc["both"] > 0 and body_jc["other"] == 0,
              f"текст не по ширине: {body_jc}")

        # --- таблицы: неразрывные строки и повтор шапки ---
        # (таблицы титульного листа — служебные, их не проверяем)
        body = doc.element.body
        kids = list(body.iterchildren())
        head_idx = next((i for i, k in enumerate(kids)
                         if k.tag.endswith("}p")
                         and ('w:pStyle w:val="Heading1"' in k.xml
                              or re.match(r"^1\. \S",
                                          "".join(k.itertext()).strip()))), None)
        tbl_i = -1
        for i, k in enumerate(kids):
            if k.tag.endswith("}tbl"):
                tbl_i += 1
                if head_idx is not None and i < head_idx:
                    continue  # таблица титульного листа
                tbl = doc.tables[tbl_i]
                rows = tbl.rows
                if not rows:
                    continue
                trpr = rows[0]._tr.find(qn("w:trPr"))
                has_hdr = trpr is not None and \
                    trpr.find(qn("w:tblHeader")) is not None
                check(has_hdr,
                      f"таблица {tbl_i}: шапка не повторяется на новой странице")
                for ri, row in enumerate(rows):
                    trpr = row._tr.find(qn("w:trPr"))
                    check(trpr is not None
                          and trpr.find(qn("w:cantSplit")) is not None,
                          f"таблица {tbl_i}, строка {ri}: нет cantSplit")

        # --- нумерованные списки: ручная нумерация, без автосписков ---
        auto = [p.text[:40] for p in doc.paragraphs
                if p.style.name in ("List Number", "List Bullet")]
        check(not auto, f"остались авто-списки (нумерация продолжится): {auto[:3]}")

        # --- подписи к рисункам центрированы ---
        for p in doc.paragraphs:
            if re.match(r"\*{0,2}Рисунок \d", p.text.strip()):
                check(str(p.alignment) == "CENTER (1)",
                      f"подпись не по центру: {p.text[:40]!r}")

        # --- рисунки не отрываются от подписи ---
        for i, p in enumerate(doc.paragraphs):
            if p._p.findall(".//" + qn("w:drawing")):
                ppr = p._p.find(qn("w:pPr"))
                check(ppr is not None and ppr.find(qn("w:keepNext")) is not None,
                      "рисунок может оторваться от подписи (нет keepNext)")

    print()
    print("ИТОГ:", "ВСЕ ПРОВЕРКИ ПРОЙДЕНЫ" if ok_all else "ЕСТЬ ЗАМЕЧАНИЯ")
    sys.exit(0 if ok_all else 1)


if __name__ == "__main__":
    main()
